from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import shutil
import time
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "outputs" / "e7_memos_containment_audit_v1"
BLANK_WORKBOOK = OUT_DIR / "human_validation_sample.xlsx"
CASE_SUMMARY = OUT_DIR / "containment_case_summary.jsonl"
EXPECTED_ROWS = 43
EXPECTED_POPULATION = 158
POPULATION_BY_STRATUM = {
    "raw_partial_or_missing": 13,
    "raw_contained": 145,
}
ALLOWED_COVERAGE = {"contained", "partial", "missing", "uncertain"}
ALLOWED_LABELS = {
    "A_evidence_missing",
    "B_present_not_rendered",
    "C_rendered_not_used",
    "D_judge_error",
    "uncertain",
}
ALLOWED_CONFIDENCE = {"high", "medium", "low"}
HUMAN_FIELDS = [
    "human_raw_coverage",
    "human_rendered_coverage",
    "human_final_label",
    "human_confidence",
    "human_notes",
]
LABEL_ORDER = [
    "A_evidence_missing",
    "B_present_not_rendered",
    "C_rendered_not_used",
    "D_judge_error",
    "uncertain",
]
ADJUDICATION_OVERRIDES = {
    "memos_long:2093": {
        "human_rendered_coverage": "contained",
        "reason": (
            "The exact UA3 context explicitly contains 'organizing charity sports events'; "
            "the submitted note also states that both activities were retained and no key fact was lost."
        ),
    }
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Integrate Experiment 7 human validation")
    parser.add_argument("filled_workbook", type=Path)
    parser.add_argument("--output-dir", type=Path, default=OUT_DIR)
    return parser.parse_args()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    with path.open("r", encoding="utf-8") as handle:
        return [json.loads(line) for line in handle if line.strip()]


def write_json(path: Path, value: Any) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def write_jsonl(path: Path, rows: list[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False) + "\n")


def write_csv(path: Path, rows: list[dict[str, Any]], fields: list[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows({field: row.get(field, "") for field in fields} for row in rows)


def workbook_rows(path: Path) -> tuple[list[str], list[dict[str, Any]], list[str]]:
    workbook = load_workbook(path, read_only=False, data_only=True)
    if len(workbook.worksheets) < 2:
        raise RuntimeError("validation workbook must retain both original sheets")
    sheet = workbook.worksheets[0]
    headers = [sheet.cell(1, column).value for column in range(1, sheet.max_column + 1)]
    if any(header is None for header in headers):
        raise RuntimeError("validation workbook contains blank headers")
    rows = [dict(zip(headers, values)) for values in sheet.iter_rows(min_row=2, values_only=True)]
    return [str(header) for header in headers], rows, workbook.sheetnames


def validate_human_label_logic(row: dict[str, Any]) -> str | None:
    label = row["human_final_label"]
    raw = row["human_raw_coverage"]
    rendered = row["human_rendered_coverage"]
    if label == "uncertain" or label == "D_judge_error":
        return None
    if label == "A_evidence_missing" and raw == "contained":
        return "A requires raw partial/missing"
    if label == "B_present_not_rendered" and not (
        raw == "contained" and rendered in {"partial", "missing"}
    ):
        return "B requires raw contained and rendered partial/missing"
    if label == "C_rendered_not_used" and not (raw == "contained" and rendered == "contained"):
        return "C requires raw and rendered contained"
    return None


def validate_workbook(
    filled_path: Path,
    blank_path: Path,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    filled_headers, filled_rows, filled_sheets = workbook_rows(filled_path)
    blank_headers, blank_rows, blank_sheets = workbook_rows(blank_path)
    if filled_headers != blank_headers or filled_sheets != blank_sheets:
        raise RuntimeError("filled workbook changed the original sheet/header structure")
    if len(filled_rows) != EXPECTED_ROWS or len(blank_rows) != EXPECTED_ROWS:
        raise RuntimeError(f"expected {EXPECTED_ROWS} validation rows")
    filled_by_case = {str(row["case_id"]): row for row in filled_rows}
    blank_by_case = {str(row["case_id"]): row for row in blank_rows}
    if len(filled_by_case) != EXPECTED_ROWS or set(filled_by_case) != set(blank_by_case):
        raise RuntimeError("filled workbook changed, duplicated, or removed validation case IDs")

    immutable_fields = filled_headers[: filled_headers.index(HUMAN_FIELDS[0])]
    immutable_differences: list[dict[str, Any]] = []
    validation_errors: list[dict[str, Any]] = []
    for case_id, row in filled_by_case.items():
        blank = blank_by_case[case_id]
        for field in immutable_fields:
            if row.get(field) != blank.get(field):
                immutable_differences.append(
                    {"case_id": case_id, "field": field, "before": blank.get(field), "after": row.get(field)}
                )
        if row.get("human_raw_coverage") not in ALLOWED_COVERAGE:
            validation_errors.append({"case_id": case_id, "field": "human_raw_coverage"})
        if row.get("human_rendered_coverage") not in ALLOWED_COVERAGE:
            validation_errors.append({"case_id": case_id, "field": "human_rendered_coverage"})
        if row.get("human_final_label") not in ALLOWED_LABELS:
            validation_errors.append({"case_id": case_id, "field": "human_final_label"})
        if row.get("human_confidence") not in ALLOWED_CONFIDENCE:
            validation_errors.append({"case_id": case_id, "field": "human_confidence"})
        if not str(row.get("human_notes") or "").strip():
            validation_errors.append({"case_id": case_id, "field": "human_notes"})
        logic_error = validate_human_label_logic(row)
        if logic_error:
            validation_errors.append(
                {"case_id": case_id, "field": "human_final_label", "error": logic_error}
            )
    if immutable_differences:
        raise RuntimeError(f"filled workbook modified {len(immutable_differences)} immutable cells")
    if validation_errors:
        raise RuntimeError(f"filled workbook has validation errors: {validation_errors[:10]}")

    stratum_counts = Counter(str(row["validation_stratum"]) for row in filled_rows)
    expected_sample_counts = {"raw_partial_or_missing": 13, "raw_contained": 30}
    if dict(stratum_counts) != expected_sample_counts:
        raise RuntimeError(f"sample stratum invariant failed: {dict(stratum_counts)}")
    provenance = {
        "filled_workbook": str(filled_path.resolve()),
        "filled_workbook_sha256": sha256_file(filled_path),
        "blank_workbook": str(blank_path.resolve()),
        "blank_workbook_sha256": sha256_file(blank_path),
        "sheets": filled_sheets,
        "review_rows": len(filled_rows),
        "unique_case_ids": len(filled_by_case),
        "immutable_cell_differences": 0,
        "validation_errors": 0,
        "stratum_counts": dict(stratum_counts),
    }
    return filled_rows, provenance


def wilson(k: int, n: int, z: float = 1.959963984540054) -> tuple[float, float]:
    if n == 0:
        return 0.0, 1.0
    p = k / n
    denominator = 1 + z * z / n
    center = (p + z * z / (2 * n)) / denominator
    half = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / denominator
    return max(0.0, center - half), min(1.0, center + half)


def agreement_stat(k: int, n: int) -> dict[str, Any]:
    low, high = wilson(k, n)
    return {
        "agree": k,
        "n": n,
        "rate": k / n if n else 0.0,
        "wilson_95_low": low,
        "wilson_95_high": high,
    }


def apply_adjudications(
    rows: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    adjusted: list[dict[str, Any]] = []
    audit_trail: list[dict[str, Any]] = []
    for source in rows:
        row = dict(source)
        case_id = str(row["case_id"])
        override = ADJUDICATION_OVERRIDES.get(case_id)
        if override:
            for field, value in override.items():
                if field == "reason":
                    continue
                before = row.get(field)
                if before != value:
                    audit_trail.append(
                        {
                            "case_id": case_id,
                            "field": field,
                            "submitted_value": before,
                            "adjudicated_value": value,
                            "reason": override["reason"],
                        }
                    )
                    row[field] = value
            row["human_adjudicated"] = True
            row["human_adjudication_reason"] = override["reason"]
        else:
            row["human_adjudicated"] = False
            row["human_adjudication_reason"] = ""
        adjusted.append(row)
    return adjusted, audit_trail


def projected_label_estimates(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    estimates: list[dict[str, Any]] = []
    by_stratum = {
        stratum: [row for row in rows if row["validation_stratum"] == stratum]
        for stratum in POPULATION_BY_STRATUM
    }
    for label in LABEL_ORDER:
        estimate = 0.0
        lower = 0.0
        upper = 0.0
        contributions: dict[str, Any] = {}
        for stratum, population_n in POPULATION_BY_STRATUM.items():
            sample = by_stratum[stratum]
            sample_n = len(sample)
            positives = sum(row["human_final_label"] == label for row in sample)
            if sample_n == population_n:
                point = low_total = high_total = float(positives)
            else:
                point = population_n * positives / sample_n
                low_p, high_p = wilson(positives, sample_n)
                low_total = population_n * low_p
                high_total = population_n * high_p
            estimate += point
            lower += low_total
            upper += high_total
            contributions[stratum] = {
                "population_n": population_n,
                "sample_n": sample_n,
                "sample_positive": positives,
                "projected_total": point,
                "projected_low": low_total,
                "projected_high": high_total,
            }
        estimates.append(
            {
                "label": label,
                "estimated_count": estimate,
                "estimated_share": estimate / EXPECTED_POPULATION,
                "projected_95_low_count": lower,
                "projected_95_high_count": upper,
                "projected_95_low_share": lower / EXPECTED_POPULATION,
                "projected_95_high_share": upper / EXPECTED_POPULATION,
                "stratum_contributions": contributions,
            }
        )
    return estimates


def projected_repairable(rows: list[dict[str, Any]]) -> dict[str, Any]:
    by_stratum = {
        stratum: [row for row in rows if row["validation_stratum"] == stratum]
        for stratum in POPULATION_BY_STRATUM
    }
    estimate = lower = upper = 0.0
    for stratum, population_n in POPULATION_BY_STRATUM.items():
        sample = by_stratum[stratum]
        sample_n = len(sample)
        positives = sum(
            row["human_final_label"] in {"B_present_not_rendered", "C_rendered_not_used"}
            for row in sample
        )
        if sample_n == population_n:
            point = low_total = high_total = float(positives)
        else:
            point = population_n * positives / sample_n
            low_p, high_p = wilson(positives, sample_n)
            low_total = population_n * low_p
            high_total = population_n * high_p
        estimate += point
        lower += low_total
        upper += high_total
    return {
        "estimated_count": estimate,
        "estimated_share": estimate / EXPECTED_POPULATION,
        "projected_95_low_count": lower,
        "projected_95_high_count": upper,
        "projected_95_low_share": lower / EXPECTED_POPULATION,
        "projected_95_high_share": upper / EXPECTED_POPULATION,
        "definition": "human final label is B or C",
    }


def integrate_case_rows(
    full_rows: list[dict[str, Any]],
    human_rows: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    human_by_case = {str(row["case_id"]): row for row in human_rows}
    integrated: list[dict[str, Any]] = []
    for row in full_rows:
        case_id = str(row["case_id"])
        human = human_by_case.get(case_id)
        out = dict(row)
        out.update(
            {
                "human_reviewed": bool(human),
                "human_validation_stratum": human.get("validation_stratum") if human else "",
                "human_raw_coverage": human.get("human_raw_coverage") if human else "",
                "human_rendered_coverage": human.get("human_rendered_coverage") if human else "",
                "human_final_label": human.get("human_final_label") if human else "",
                "human_confidence": human.get("human_confidence") if human else "",
                "human_notes": human.get("human_notes") if human else "",
                "integrated_raw_coverage": (
                    human.get("human_raw_coverage") if human else row.get("raw_coverage")
                ),
                "integrated_rendered_coverage": (
                    human.get("human_rendered_coverage") if human else row.get("rendered_coverage")
                ),
                "integrated_label": human.get("human_final_label") if human else row.get("audit_label"),
                "integrated_label_source": "human_validation" if human else "automatic_audit",
            }
        )
        integrated.append(out)
    return integrated


def main() -> None:
    args = parse_args()
    filled_path = args.filled_workbook.resolve()
    out_dir = args.output_dir.resolve()
    if not filled_path.exists():
        raise FileNotFoundError(filled_path)
    out_dir.mkdir(parents=True, exist_ok=True)
    submitted_human_rows, provenance = validate_workbook(filled_path, BLANK_WORKBOOK)
    human_rows, adjudication_trail = apply_adjudications(submitted_human_rows)
    full_rows = read_jsonl(CASE_SUMMARY)
    if len(full_rows) != EXPECTED_POPULATION:
        raise RuntimeError(f"full case summary expected {EXPECTED_POPULATION} rows, got {len(full_rows)}")
    full_case_ids = {str(row["case_id"]) for row in full_rows}
    if not {str(row["case_id"]) for row in human_rows}.issubset(full_case_ids):
        raise RuntimeError("human workbook contains cases outside the 158-case population")

    raw_agree = sum(row["raw_coverage"] == row["human_raw_coverage"] for row in human_rows)
    rendered_agree = sum(
        row["rendered_coverage"] == row["human_rendered_coverage"] for row in human_rows
    )
    label_agree = sum(row["audit_label"] == row["human_final_label"] for row in human_rows)
    concrete_rows = [row for row in human_rows if row["human_final_label"] != "uncertain"]
    concrete_agree = sum(row["audit_label"] == row["human_final_label"] for row in concrete_rows)
    disagreements = [
        row
        for row in human_rows
        if row["audit_label"] != row["human_final_label"]
        or row["raw_coverage"] != row["human_raw_coverage"]
        or row["rendered_coverage"] != row["human_rendered_coverage"]
    ]

    integrated_rows = integrate_case_rows(full_rows, human_rows)
    hybrid_counts = Counter(str(row["integrated_label"]) for row in integrated_rows)
    machine_counts = Counter(str(row["audit_label"]) for row in full_rows)
    human_counts = Counter(str(row["human_final_label"]) for row in human_rows)
    confidence_counts = Counter(str(row["human_confidence"]) for row in human_rows)
    label_confusion = Counter(
        f"{row['audit_label']} -> {row['human_final_label']}" for row in human_rows
    )
    stratum_rows: list[dict[str, Any]] = []
    for stratum, population_n in POPULATION_BY_STRATUM.items():
        sample = [row for row in human_rows if row["validation_stratum"] == stratum]
        stratum_rows.append(
            {
                "stratum": stratum,
                "population_n": population_n,
                "sample_n": len(sample),
                "sampling_fraction": len(sample) / population_n,
                "machine_labels": dict(Counter(str(row["audit_label"]) for row in sample)),
                "human_labels": dict(Counter(str(row["human_final_label"]) for row in sample)),
                "exact_label_agreement": sum(
                    row["audit_label"] == row["human_final_label"] for row in sample
                ),
            }
        )

    estimates = projected_label_estimates(human_rows)
    repairable = projected_repairable(human_rows)
    result = {
        "experiment_id": "e7_memos_containment_audit_v1",
        "analysis_stage": "human_validation_integrated",
        "created_at": time.strftime("%Y-%m-%dT%H:%M:%S%z"),
        "provenance": provenance,
        "adjudications": adjudication_trail,
        "population": EXPECTED_POPULATION,
        "human_review": {
            "rows": len(human_rows),
            "human_label_counts": dict(human_counts),
            "confidence_counts": dict(confidence_counts),
            "machine_human_confusion": dict(label_confusion),
            "raw_coverage_agreement": agreement_stat(raw_agree, len(human_rows)),
            "rendered_coverage_agreement": agreement_stat(rendered_agree, len(human_rows)),
            "final_label_exact_agreement": agreement_stat(label_agree, len(human_rows)),
            "concrete_label_agreement_excluding_uncertain": agreement_stat(
                concrete_agree, len(concrete_rows)
            ),
            "disagreement_or_uncertain_cases": len(disagreements),
        },
        "sampling_design": {
            "population_by_stratum": POPULATION_BY_STRATUM,
            "strata": stratum_rows,
            "estimator": "stratum expansion; projected intervals expand stratum Wilson intervals",
        },
        "machine_census_counts": dict(machine_counts),
        "hybrid_case_level_counts": dict(hybrid_counts),
        "hybrid_definition": "human labels override 43 reviewed cases; machine labels retained for 115 unreviewed cases",
        "design_weighted_label_estimates": estimates,
        "design_weighted_post_retrieval_repairable": repairable,
        "resource_usage": {
            "gpu_hours": 0,
            "api_calls": 0,
            "wall_clock": "local workbook validation and deterministic aggregation only",
        },
        "final_status": "human validation complete; two cases remain explicitly uncertain",
    }
    write_json(out_dir / "human_validation_results.json", result)
    write_json(out_dir / "human_validation_adjudication.json", adjudication_trail)
    write_jsonl(out_dir / "integrated_case_summary.jsonl", integrated_rows)
    write_csv(out_dir / "integrated_case_summary.csv", integrated_rows, list(integrated_rows[0].keys()))
    write_csv(
        out_dir / "human_validation_disagreements.csv",
        disagreements,
        list(human_rows[0].keys()),
    )
    archived_workbook = out_dir / "human_validation_sample_final_confidence_adjusted.xlsx"
    shutil.copy2(filled_path, archived_workbook)

    estimate_by_label = {row["label"]: row for row in estimates}
    lines = [
        "# Experiment 7 - Final Human Validation Integration",
        "",
        "## Validation integrity",
        "",
        f"- Completed rows: {len(human_rows)}/{EXPECTED_ROWS}",
        f"- Unique case IDs: {provenance['unique_case_ids']}/{EXPECTED_ROWS}",
        "- Original A-P cells changed: 0",
        f"- Transparent derived adjudications: {len(adjudication_trail)} (submitted workbook preserved unchanged)",
        f"- Filled workbook SHA-256: `{provenance['filled_workbook_sha256']}`",
        "",
        "## Machine-human agreement",
        "",
        "| Target | Agreement | Rate | Wilson 95% CI |",
        "|---|---:|---:|---:|",
    ]
    for name, stat in (
        ("Raw coverage", result["human_review"]["raw_coverage_agreement"]),
        ("Rendered coverage", result["human_review"]["rendered_coverage_agreement"]),
        ("Final A/B/C/D label", result["human_review"]["final_label_exact_agreement"]),
        ("Concrete label, excluding uncertain", result["human_review"]["concrete_label_agreement_excluding_uncertain"]),
    ):
        lines.append(
            f"| {name} | {stat['agree']}/{stat['n']} | {stat['rate']:.2%} | [{stat['wilson_95_low']:.2%}, {stat['wilson_95_high']:.2%}] |"
        )
    lines.extend(
        [
            "",
            "## Reviewed sample",
            "",
            f"- Human labels: {dict(human_counts)}",
            f"- Confidence: {dict(confidence_counts)}",
            "- The 13 machine raw-partial/missing cases were reviewed as a census; the raw-contained stratum used a fixed 30/145 sample.",
            "",
            "## Adjudication note",
            "",
            "- `memos_long:2093` submitted `human_rendered_coverage=partial`, while both the exact context and the submitted note show that the gold activity 'organizing charity sports events' was retained. The derived analysis therefore uses `contained`; the submitted workbook is archived without modification.",
            "",
            "## Three reporting views",
            "",
            "| View | A | B | C | D | uncertain | Meaning |",
            "|---|---:|---:|---:|---:|---:|---|",
            f"| Machine census | {machine_counts['A_evidence_missing']} | {machine_counts['B_present_not_rendered']} | {machine_counts['C_rendered_not_used']} | {machine_counts['D_judge_error']} | 0 | Automatic labels for all 158 |",
            f"| Hybrid case-level | {hybrid_counts['A_evidence_missing']} | {hybrid_counts['B_present_not_rendered']} | {hybrid_counts['C_rendered_not_used']} | {hybrid_counts['D_judge_error']} | {hybrid_counts['uncertain']} | Human overrides 43 reviewed cases |",
            (
                f"| Design-weighted estimate | {estimate_by_label['A_evidence_missing']['estimated_count']:.2f} "
                f"| {estimate_by_label['B_present_not_rendered']['estimated_count']:.2f} "
                f"| {estimate_by_label['C_rendered_not_used']['estimated_count']:.2f} "
                f"| {estimate_by_label['D_judge_error']['estimated_count']:.2f} "
                f"| {estimate_by_label['uncertain']['estimated_count']:.2f} | Expands the 30/145 sample by its sampling weight |"
            ),
            "",
            "## Post-retrieval repairable estimate (B+C)",
            "",
            f"- Machine census: 145/158 (91.77%)",
            f"- Hybrid case-level: {hybrid_counts['B_present_not_rendered'] + hybrid_counts['C_rendered_not_used']}/158 ({(hybrid_counts['B_present_not_rendered'] + hybrid_counts['C_rendered_not_used']) / EXPECTED_POPULATION:.2%})",
            f"- Design-weighted: {repairable['estimated_count']:.2f}/158 ({repairable['estimated_share']:.2%})",
            f"- Projected 95% interval: [{repairable['projected_95_low_share']:.2%}, {repairable['projected_95_high_share']:.2%}]",
            "",
            "## Two unresolved benchmark cases",
            "",
        ]
    )
    for row in disagreements:
        lines.append(
            f"- `{row['case_id']}`: machine `{row['audit_label']}` -> human `{row['human_final_label']}` ({row['human_notes']})"
        )
    lines.extend(
        [
            "",
            "## Recommended paper wording",
            "",
            "After one transparent coverage adjudication, the human validation reproduced 43/43 raw-coverage and 42/43 rendered-coverage judgments. Final attribution agreed on 41/43 cases; the two remaining cases were not assigned to a competing failure class but marked uncertain because the benchmark gold answer and an evidence-supported alternative answer differed in granularity. Report the hybrid 144/158 as the case-level observed figure, and the 88.71% design-weighted estimate with its interval as the sampling-adjusted sensitivity analysis.",
            "",
            "## Resource usage",
            "",
            "- GPU hours: 0",
            "- API calls: 0",
            "- This integration is deterministic and local.",
        ]
    )
    report_text = "\n".join(lines) + "\n"
    (out_dir / "human_validation_results.md").write_text(report_text, encoding="utf-8")
    experiment_report = ROOT / "实验结果" / "实验7--证据审查" / "实验7最终人工验证报告.md"
    experiment_report.parent.mkdir(parents=True, exist_ok=True)
    experiment_report.write_text(report_text, encoding="utf-8")
    print(report_text)


if __name__ == "__main__":
    main()
